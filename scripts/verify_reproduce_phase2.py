import sys
import os
import json
import hashlib
import sqlite3
from pathlib import Path
import pandas as pd
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

PROJECT_ROOT = Path("C:/Users/Arindam Das/Desktop/project/MIP_Android")
sys.path.insert(0, str(PROJECT_ROOT))

from core.config import MIPScannerConfig

report = {}

# 1. Recompute Strategy Fingerprint
frozen_params = {
    "ema_regime_span": 20,
    "ema_rs_span": 200,
    "ema_trend_span": 200,
    "exit_buffer_rank": 40,
    "max_per_sector": 2,
    "momentum_lookback_bars": 252,
    "retrace_threshold": 0.80,
    "top_n_candidates": 20,
    "volatility_floor": 0.05,
}
serialized = json.dumps(frozen_params, sort_keys=True)
computed_fp = hashlib.sha256(serialized.encode("utf-8")).hexdigest()
canonical_fp = "9036fae5c41a27713901c8e82a97dfa07bc2f26a86f8e3231fdb8868d78673de"

report["strategy_fingerprint"] = {
    "canonical": canonical_fp,
    "computed": computed_fp,
    "match": computed_fp == canonical_fp
}

# 2. Check 9 frozen scalar parameters in MIPScannerConfig
cfg = MIPScannerConfig()
config_params_match = True
param_checks = {}
for k, v in frozen_params.items():
    actual_v = getattr(cfg, k, None)
    match = actual_v == v
    param_checks[k] = {"expected": v, "actual": actual_v, "match": match}
    if not match:
        config_params_match = False

report["frozen_parameters"] = {
    "all_matched": config_params_match,
    "details": param_checks
}

# 3. Recompute configuration hash
computed_cfg_hash = cfg.get_config_hash()
report["config_hash"] = {
    "computed": computed_cfg_hash,
    "matches_checkpoint": computed_cfg_hash == "6c8981f97b010ad4fd3225bd7553022dba84c080e4139ebd3b82b2eb1677e239"
}

# 4. Verify Canonical Volar Formula
def check_volar_formula():
    with open(PROJECT_ROOT / "core" / "scanner.py", "r", encoding="utf-8") as f:
        scanner_src = f.read()
    with open(PROJECT_ROOT / "core" / "backtest.py", "r", encoding="utf-8") as f:
        backtest_src = f.read()
        
    has_scanner_volar = 'bars["volar_score"] = bars["ret_252"] / bars["vol_252"].clip(lower=self.config.volatility_floor)' in scanner_src
    has_backtest_volar = 'ret_1y / vol_252.clip(lower=0.05)' in backtest_src
    return {
        "scanner_has_canonical_volar": has_scanner_volar,
        "backtest_has_canonical_volar": has_backtest_volar,
        "valid": has_scanner_volar and has_backtest_volar
    }

report["volar_formula"] = check_volar_formula()

# 5. Verify Provenance Manifest
def check_provenance():
    manifest_path = PROJECT_ROOT / "certificates" / "PROVENANCE_MANIFEST.json"
    with open(manifest_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    modules = data.get("modules", [])
    results = []
    all_match = True
    for m in modules:
        fname = m["filename"]
        exp_hash = m["sha256_migrated"]
        fpath = PROJECT_ROOT / "core" / fname
        if not fpath.exists():
            results.append({"filename": fname, "status": "MISSING"})
            all_match = False
            continue
        h = hashlib.sha256()
        with open(fpath, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                h.update(chunk)
        act_hash = h.hexdigest()
        m_match = act_hash == exp_hash
        results.append({
            "filename": fname,
            "expected_hash": exp_hash,
            "actual_hash": act_hash,
            "match": m_match
        })
        if not m_match:
            all_match = False
            
    return {
        "total_modules": len(modules),
        "all_match": all_match,
        "modules": results
    }

report["provenance_manifest"] = check_provenance()

# 6. Checkpoints Audit
def check_checkpoints():
    ckpt_dir = PROJECT_ROOT / "artifacts" / "phase-ranking-repair" / "state"
    gates = []
    all_valid = True
    for i in range(1, 11):
        cpath = ckpt_dir / f"gate-{i}.json"
        if not cpath.exists():
            gates.append({"gate": f"gate-{i}", "status": "MISSING"})
            all_valid = False
            continue
        with open(cpath, "r", encoding="utf-8") as f:
            cdata = json.load(f)
        gates.append({
            "gate": cdata.get("gate"),
            "status": cdata.get("status"),
            "timestamp": cdata.get("timestamp"),
            "has_code_hash": "code_hash" in cdata,
            "has_dataset_fp": "dataset_fingerprint" in cdata,
            "has_config_hash": "configuration_hash" in cdata or "config_hash" in cdata
        })
    return {
        "all_present_and_valid": all_valid,
        "gates": gates
    }

report["checkpoints"] = check_checkpoints()

# 7. Universe Reconciliation
def check_universe():
    total_mkt_csv = PROJECT_ROOT / "data" / "raw_reference" / "ind_niftytotalmarket_list.csv"
    db_path = PROJECT_ROOT / "data" / "universe.db"
    
    df_csv = pd.read_csv(total_mkt_csv)
    csv_syms = set(df_csv["Symbol"].str.strip().str.upper())
    
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    db_syms = set(r[0] for r in cur.execute("SELECT DISTINCT symbol FROM prices WHERE date='2026-10-09'").fetchall())
    conn.close()
    
    missing_in_db = sorted(list(csv_syms - db_syms))
    
    # 8 exclusions documented
    expected_exclusions = sorted(['BAGMANE', 'BIRET', 'EMBASSY', 'DUMMYHEG', 'DUMMYINGL1', 'DUMMYINGL2', 'DUMMYINXGN', 'TRIVENIPT'])
    
    return {
        "total_mkt_csv_count": len(csv_syms),
        "db_matched_count": len(db_syms.intersection(csv_syms)),
        "db_total_symbols_20261009": len(db_syms),
        "missing_in_db": missing_in_db,
        "missing_matches_expected": missing_in_db == expected_exclusions
    }

report["universe_reconciliation"] = check_universe()

# 8. Sector Map Audit
def check_sector_map():
    sec_map_path = PROJECT_ROOT / "data" / "symbol_sector_map.json"
    with open(sec_map_path, "r", encoding="utf-8") as f:
        smap = json.load(f)
    unique_secs = sorted(list(set(smap.values())))
    has_infra_media = "INFRA_MEDIA" in unique_secs
    return {
        "total_symbols": len(smap),
        "unique_sectors_count": len(unique_secs),
        "has_legacy_infra_media": has_infra_media,
        "sectors": unique_secs
    }

report["sector_map"] = check_sector_map()

# 9. Deliverables Inspection
def check_deliverables():
    wb_path = PROJECT_ROOT / "reports" / "MIP1_Momentum_Scanner_2026-10-09_NIFTY750TOTALMARKET.xlsx"
    csv_path = PROJECT_ROOT / "reports" / "screener_output_live.csv"
    
    wb_exists = wb_path.exists()
    csv_exists = csv_path.exists()
    
    wb_info = {}
    if wb_exists:
        wb = openpyxl.load_workbook(wb_path, data_only=True)
        wb_info = {
            "size_bytes": os.path.getsize(wb_path),
            "sheet_count": len(wb.sheetnames),
            "sheets": wb.sheetnames
        }
        wb.close()
        
    csv_info = {}
    if csv_exists:
        df = pd.read_csv(csv_path)
        csv_info = {
            "rows": len(df),
            "columns": list(df.columns)
        }
        
    return {
        "workbook_exists": wb_exists,
        "workbook_info": wb_info,
        "screener_csv_exists": csv_exists,
        "screener_csv_info": csv_info
    }

report["deliverables"] = check_deliverables()

# 10. Artifact Registry Inspection
def check_artifact_registry():
    reg_path = PROJECT_ROOT / "artifacts" / "ARTIFACT_REGISTRY.md"
    exists = reg_path.exists()
    count = 0
    if exists:
        with open(reg_path, "r", encoding="utf-8") as f:
            lines = f.readlines()
        count = sum(1 for line in lines if line.strip().startswith("|") and not line.strip().startswith("| #") and not line.strip().startswith("|---"))
    return {
        "registry_exists": exists,
        "registered_count": count,
        "within_cap": count <= 10
    }

report["artifact_registry"] = check_artifact_registry()

print(json.dumps(report, indent=2))
